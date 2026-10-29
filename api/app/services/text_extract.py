"""Text extraction from resumes and JD attachments (PDF, DOCX, TXT)."""
from __future__ import annotations

import io
from pathlib import PurePath

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".rtf"}


def extract_text(file_name: str, data: bytes) -> str:
    ext = PurePath(file_name).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if ext == ".docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if ext in {".txt", ".md", ".rtf"}:
        return data.decode("utf-8", errors="replace").strip()
    raise ValueError(f"Unsupported file type: {ext or file_name}")
