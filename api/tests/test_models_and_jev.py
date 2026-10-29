"""Model routing per task, fallback to Claude, Jev triage + safety net, settings rules, multi-contacts, JD import."""
import asyncio
import io
import json

import docx
import pytest
from conftest import TENANT_A, TENANT_B
from test_email_agent import JOB, cls, fresh_mail, msg, run_scan  # noqa: F401 — fixtures/helpers

from app.agent import email_agent
from app.llm import FakeProvider
from app.llm.providers import FallbackProvider
from app.llm.tenant import route_for
from app.services import jev
from app.services.graph import Attachment

A = f"/tenants/{TENANT_A}"


# ── routing + fallback ──────────────────────────────────────────────────────

def test_route_for_uses_task_then_default():
    st = {"llm_provider": "anthropic", "llm_model": "claude-sonnet-5-5", "llm_base_url": None,
          "model_routes": {"default": {"provider": "openrouter", "model": "qwen/qwen-2.5-72b-instruct"},
                           "resume_parsing": {"provider": "openai_compatible", "model": "qwen2.5:14b",
                                              "base_url": "http://host.docker.internal:11434/v1", "fallback": True}}}
    assert route_for(st, "resume_parsing")["model"] == "qwen2.5:14b"
    assert route_for(st, "email_extraction")["provider"] == "openrouter"      # falls to workspace default
    assert route_for({**st, "model_routes": {}}, "email_triage")["model"] == "claude-sonnet-5-5"


def test_fallback_answers_when_cheap_model_returns_junk():
    cheap, claude = FakeProvider(["Sure! here is the info, no json"]), FakeProvider(['{"first_name": "Ana"}'])
    fb = FallbackProvider(cheap, claude)
    data, _ = asyncio.run(fb.complete_json("s", "u"))
    assert data == {"first_name": "Ana"} and fb.last_used is claude


# ── settings rules ──────────────────────────────────────────────────────────

def test_settings_accept_routes_and_jev_and_validate(api_client):
    c = api_client
    ok = c.patch(f"{A}/settings", json={
        "model_routes": {"email_extraction": {"provider": "openrouter", "model": "qwen/qwen-2.5-72b-instruct", "fallback": True},
                         "resume_parsing": None},
        "triage_provider": "jev", "jev_model": "jev-1.13.0", "review_below_confidence": 0.7,
        "jev_config": {"categories": {"resume": "A CV or resume"}, "body_chars": 4000}})
    assert ok.status_code == 200, ok.text
    s = ok.json()
    assert s["model_routes"]["email_extraction"]["fallback"] is True and "resume_parsing" not in s["model_routes"]
    assert s["jev_config"] == {"categories": {"resume": "A CV or resume"}, "body_chars": 4000}
    assert c.patch(f"{A}/settings", json={"model_routes": {"made_up_task": {"provider": "anthropic", "model": "x"}}}).status_code == 422
    assert c.patch(f"{A}/settings", json={"jev_config": {"categories": {"spam": "x"}}}).status_code == 422
    assert c.patch(f"{A}/settings", json={"model_routes": {"default": {"provider": "gpt", "model": "x"}}}).status_code == 422
    c.patch(f"{A}/settings", json={"triage_provider": "llm", "model_routes": {}, "jev_config": {}})


def test_ai_status_reports_keys_without_revealing_them(api_client):
    s = api_client.get(f"{A}/ai/status").json()
    assert set(s["keys"]) == {"anthropic", "openrouter", "typesafe"} and all(isinstance(v, bool) for v in s["keys"].values())
    assert s["can_edit"] and "email_extraction" in s["tasks"] and s["jev_defaults"]["categories"]["resume"]


# ── Jev triage ──────────────────────────────────────────────────────────────

class FakeJev:
    def __init__(self, probs: dict[str, float] | Exception):
        self.probs, self.requests = probs, []

    async def classify(self, request):
        self.requests.append(request)
        if isinstance(self.probs, Exception):
            raise self.probs
        choice = max(self.probs, key=self.probs.get)
        return jev.Triage(choice, self.probs[choice], self.probs, "jev-1.13.0")


@pytest.fixture
def jev_on(api_client, monkeypatch):
    api_client.patch(f"{A}/settings", json={"triage_provider": "jev", "review_below_confidence": 0.6})
    yield monkeypatch
    api_client.patch(f"{A}/settings", json={"triage_provider": "llm"})


def test_jev_sorts_email_and_claude_extracts(api_client, fresh_mail, tmp_path, jev_on):
    fake = FakeJev({"job_requisition": 0.91, "resume": 0.03, "interview_feedback": 0.01, "client_correspondence": 0.03,
                    "candidate_reply": 0.01, "other": 0.01})
    jev_on.setattr(email_agent, "jev_client", lambda: fake)
    r, llm = run_scan(fresh_mail, tmp_path, [msg(60, "Req: Willow analyst", "Need a Willow analyst")], [JOB])
    assert r.queued == 1
    assert len(llm.calls) == 1   # the LLM only extracted; Jev did the sorting
    q = fake.requests[0]["questions"]["email_type"]
    assert q["type"] == "choice" and set(q["criteria"]) == set(jev.CLASSES)
    assert fake.requests[0]["state"]["subject"] == "Req: Willow analyst"
    recent = api_client.get(f"{A}/jev/recent").json()
    assert recent[0]["triage"]["provider"] == "jev" and recent[0]["triage"]["probabilities"]["job_requisition"] == 0.91


def test_jev_unsure_not_recruiting_still_goes_to_review(api_client, fresh_mail, tmp_path, jev_on):
    fake = FakeJev({"job_requisition": 0.05, "resume": 0.38, "interview_feedback": 0.02, "client_correspondence": 0.03,
                    "candidate_reply": 0.02, "other": 0.50})
    jev_on.setattr(email_agent, "jev_client", lambda: fake)
    resume = json.dumps({"first_name": "Noor", "last_name": "Aziz", "email": "noor@example.com", "confidence": 0.8})
    r, _ = run_scan(fresh_mail, tmp_path, [msg(61, "fwd", "see attached cv", minutes_ago=5)], [resume])
    assert r.queued == 1 and r.ignored == 0
    t = api_client.get(f"{A}/jev/recent").json()[0]["triage"]
    assert t["unsure"]["was"] == "other"


def test_jev_down_falls_back_to_llm(api_client, fresh_mail, tmp_path, jev_on):
    jev_on.setattr(email_agent, "jev_client", lambda: FakeJev(jev.JevError("TypeSafe rejected the API key")))
    r, llm = run_scan(fresh_mail, tmp_path, [msg(62, "Newsletter", "hi", minutes_ago=4)], [cls("other")])
    assert r.ignored == 1 and len(llm.calls) == 1
    t = api_client.get(f"{A}/jev/recent").json()[0]["triage"]
    assert t["provider"] == "fake" and "rejected" in t["jev_error"]


def test_jev_request_shape_and_corrections():
    cfg = jev.merged_config({"categories": {"resume": "Custom resume text", "bogus": "x"}, "body_chars": 600})
    assert cfg["categories"]["resume"] == "Custom resume text" and "bogus" not in cfg["categories"]
    req = jev.build_request(cfg, "jev-latest", "S", "A <a@b.c>", "x" * 5000, [("JD.pdf", "y" * 5000)],
                            ["email from a@b.c was labelled other, but it was actually resume"])
    assert len(req["state"]["body"]) == 600 and len(req["state"]["attachments"][0]["excerpt"]) == 1500
    assert "actually resume" in req["questions"]["email_type"]["instructions"]


# ── JD attachments + job page import ───────────────────────────────────────

def jd_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Epic Willow Pharmacy Analyst", 1)
    d.add_heading("Responsibilities", 2)
    d.add_paragraph("Build Willow orderables", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_email_jd_attachment_becomes_the_description(api_client, fresh_mail, tmp_path):
    att = Attachment("FH Willow JD.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", jd_docx())
    job = json.loads(JOB)
    job["description"] = "A short summary the model wrote"
    run_scan(fresh_mail, tmp_path, [msg(70, "New req", "JD attached", atts=[att], minutes_ago=2)],
             [cls("job_requisition"), json.dumps(job)])
    item = api_client.get(f"{A}/inbox").json()["items"][0]
    d = item["extracted"]["description"]
    assert d.startswith("# Epic Willow Pharmacy Analyst") and "- Build Willow orderables" in d
    assert item["extracted"]["_description_from"] == ["FH Willow JD.docx"]

    # The same JD can be pulled into the job later from the job page
    created = api_client.post(f"{A}/inbox/{item['id']}/approve", json={"final": {**item["extracted"], "description": ""}}).json()
    jid = created["entity_id"]
    srcs = api_client.get(f"{A}/jobs/{jid}/jd-sources").json()["attachments"]
    assert srcs[0]["file_name"] == "FH Willow JD.docx" and srcs[0]["suggested"]
    prev = api_client.post(f"{A}/jobs/{jid}/jd-preview/attachment/{srcs[0]['id']}").json()
    assert "## Responsibilities" in prev["markdown"]
    up = api_client.post(f"{A}/jobs/{jid}/jd-preview/file", files={"file": ("jd.txt", b"ABOUT THE ROLE\n\xe2\x80\xa2 Build things\n")}).json()
    assert "## ABOUT THE ROLE" in up["markdown"] and "- Build things" in up["markdown"]
    assert api_client.post(f"{A}/jobs/{jid}/jd-preview/file", files={"file": ("x.png", b"x")}).status_code == 422


# ── contacts ────────────────────────────────────────────────────────────────

def test_add_several_contacts_at_once(api_client):
    c = api_client
    cl = c.post(f"{A}/clients", json={"name": "Interior Health"}).json()
    rows = [{"first_name": "Kim", "email": "kim@interior-demo.ca", "is_hiring_manager": True, "is_primary": True},
            {"first_name": "Raj", "last_name": "Singh", "title": "Procurement", "email": "raj@interior-demo.ca"}]
    out = c.post(f"{A}/clients/{cl['id']}/contacts", json=rows)
    assert out.status_code == 201 and len(out.json()) == 2
    assert len(c.get(f"{A}/contacts", params={"client_id": cl["id"]}).json()) == 2
    # all or nothing: one duplicate → nothing saved
    again = c.post(f"{A}/clients/{cl['id']}/contacts", json=[{"first_name": "New", "email": "new@interior-demo.ca"},
                                                               {"first_name": "Kim", "email": "kim@interior-demo.ca"}])
    assert again.status_code == 409
    assert len(c.get(f"{A}/contacts", params={"client_id": cl["id"]}).json()) == 2
    assert c.post(f"/tenants/{TENANT_B}/clients/{cl['id']}/contacts", json=rows).status_code in (403, 404)


def test_contact_details_edit_and_duplicate_email(api_client):
    c = api_client
    a = c.post(f"{A}/candidates", json={"first_name": "Zed", "email": "zed@example.com"}).json()
    b = c.post(f"{A}/candidates", json={"first_name": "Yan"}).json()
    r = c.patch(f"{A}/candidates/{b['id']}", json={"phone": "604-555-0199", "linkedin_url": "https://linkedin.com/in/yan",
                                                     "location": "Burnaby, BC", "availability": "Immediately"})
    assert r.json()["phone"] == "604-555-0199"
    assert c.patch(f"{A}/candidates/{b['id']}", json={"email": "zed@example.com"}).status_code == 409
    assert c.patch(f"{A}/candidates/{a['id']}", json={"phone": None}).json()["phone"] is None
